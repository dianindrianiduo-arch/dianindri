"""Fill the exam template docx with PAI Kelas XI exam questions."""
import docx
from docx.oxml.ns import qn
from copy import deepcopy

SRC = '/root/.claude/uploads/b3a0fb48-92a3-493e-8b53-f0a69f1f7e44/ea778bad-KARTU_SOAL_UasGENAP_2526okeaslikosongan.docx'

doc = docx.Document(SRC)
body = doc.element.body


# =====================================================================
# CONTENT
# =====================================================================

NAMA_GURU = "NUR LAILY FAUZIAH, S.Pd.I."
MAPEL = "Pendidikan Agama Islam (PAI)"
KELAS = "XI (Sebelas)"
TAHUN = "2025/2026"

# Pilihan Ganda
PG = [
    # ===== BAB 1: PERNIKAHAN (5 soal) =====
    {
        "no": 1, "bab": "Pernikahan dalam Islam",
        "elemen": "Fikih – Pernikahan dalam Islam",
        "cp": "Peserta didik memahami ketentuan pernikahan dalam Islam dan hikmahnya.",
        "tp": "Peserta didik dapat menjelaskan pengertian pernikahan dalam Islam.",
        "indikator_kompetensi": "Memahami pengertian pernikahan dalam Islam.",
        "indikator_soal": "Disajikan narasi tentang dasar disyariatkannya pernikahan, peserta didik dapat menentukan pengertian nikah yang paling tepat.",
        "tingkat": "Rendah", "kunci": "B",
        "soal": (
            "Pernikahan dalam Islam memiliki landasan yang kuat di dalam Al-Qur'an. Salah satu ayat yang menjadi dasar disyariatkannya pernikahan terdapat dalam QS An-Nur ayat 32. Pengertian nikah secara istilah dalam Islam adalah ...\n"
            "A. Hubungan kerja sama dua pihak dalam masalah harta\n"
            "B. Ikatan lahir batin antara seorang laki-laki dan perempuan sebagai suami istri yang sah berdasarkan hukum Allah Swt.\n"
            "C. Perjanjian dagang antara dua keluarga\n"
            "D. Persahabatan yang erat antara dua keluarga besar\n"
            "E. Hubungan pertemanan tanpa ikatan resmi"
        ),
    },
    {
        "no": 2, "bab": "Pernikahan dalam Islam",
        "elemen": "Fikih – Pernikahan dalam Islam",
        "cp": "Peserta didik memahami hukum pernikahan berdasarkan kondisi seseorang.",
        "tp": "Peserta didik dapat menentukan hukum nikah berdasarkan kondisi tertentu.",
        "indikator_kompetensi": "Menentukan hukum nikah.",
        "indikator_soal": "Disajikan narasi tentang seseorang yang sudah mampu dan dikhawatirkan zina, peserta didik dapat menentukan hukum nikahnya.",
        "tingkat": "Sedang", "kunci": "C",
        "soal": (
            "Hukum asal nikah dalam Islam adalah mubah, namun dapat berubah sesuai kondisi seseorang. Bagi seorang muslim yang sudah mampu lahir batin dan dikhawatirkan terjerumus pada perbuatan zina jika tidak segera menikah, maka hukum nikah baginya adalah ...\n"
            "A. Sunah\n"
            "B. Mubah\n"
            "C. Wajib\n"
            "D. Makruh\n"
            "E. Haram"
        ),
    },
    {
        "no": 3, "bab": "Pernikahan dalam Islam",
        "elemen": "Fikih – Pernikahan dalam Islam",
        "cp": "Peserta didik memahami rukun pernikahan.",
        "tp": "Peserta didik dapat membedakan rukun nikah dan yang bukan rukun nikah.",
        "indikator_kompetensi": "Mengidentifikasi rukun nikah.",
        "indikator_soal": "Disajikan beberapa unsur pernikahan, peserta didik dapat menentukan yang BUKAN rukun nikah.",
        "tingkat": "Sedang", "kunci": "E",
        "soal": (
            "Pernikahan dapat dikatakan sah apabila telah memenuhi rukun-rukunnya. Berikut ini yang BUKAN termasuk rukun nikah adalah ...\n"
            "A. Calon suami\n"
            "B. Calon istri\n"
            "C. Wali nikah\n"
            "D. Dua orang saksi\n"
            "E. Mahar dengan nilai tertentu"
        ),
    },
    {
        "no": 4, "bab": "Pernikahan dalam Islam",
        "elemen": "Fikih – Pernikahan dalam Islam",
        "cp": "Peserta didik memahami urutan wali nikah.",
        "tp": "Peserta didik dapat menentukan urutan wali nikah yang paling utama.",
        "indikator_kompetensi": "Menentukan wali nikah yang paling berhak.",
        "indikator_soal": "Disajikan beberapa pilihan kerabat, peserta didik dapat menentukan wali nikah yang paling berhak menurut fikih.",
        "tingkat": "Sedang", "kunci": "C",
        "soal": (
            "Wali merupakan salah satu rukun nikah yang harus ada dalam pelaksanaan akad. Wali yang paling berhak menikahkan seorang perempuan menurut urutan dalam fikih adalah ...\n"
            "A. Paman dari pihak ibu\n"
            "B. Saudara laki-laki seibu\n"
            "C. Ayah kandung\n"
            "D. Kakek dari pihak ibu\n"
            "E. Tetangga yang dipercaya"
        ),
    },
    {
        "no": 5, "bab": "Pernikahan dalam Islam",
        "elemen": "Fikih – Pernikahan dalam Islam",
        "cp": "Peserta didik memahami hikmah pernikahan.",
        "tp": "Peserta didik dapat menjelaskan hikmah utama pernikahan.",
        "indikator_kompetensi": "Memahami hikmah pernikahan.",
        "indikator_soal": "Disajikan pernyataan tentang hikmah pernikahan, peserta didik dapat menentukan hikmah utama yang sesuai ajaran Islam.",
        "tingkat": "Sedang", "kunci": "B",
        "soal": (
            "Pernikahan memiliki hikmah yang sangat agung bagi kehidupan manusia. Salah satu hikmah pernikahan yang paling utama adalah ...\n"
            "A. Memperbanyak harta benda\n"
            "B. Menjaga kehormatan diri, melahirkan keturunan yang sah, dan membentuk keluarga sakinah mawaddah wa rahmah\n"
            "C. Mempererat hubungan bisnis antar keluarga\n"
            "D. Mendapatkan jabatan yang lebih tinggi di masyarakat\n"
            "E. Menambah pengikut di media sosial"
        ),
    },
    # ===== BAB 2: ADAB MENGGUNAKAN MEDIA SOSIAL (5 soal) =====
    {
        "no": 6, "bab": "Adab Menggunakan Media Sosial",
        "elemen": "Akhlak – Adab Menggunakan Media Sosial",
        "cp": "Peserta didik memahami adab seorang muslim dalam bermedia sosial.",
        "tp": "Peserta didik dapat menentukan sikap muslim yang tepat dalam menggunakan media sosial.",
        "indikator_kompetensi": "Memahami adab bermedia sosial.",
        "indikator_soal": "Disajikan narasi tentang penggunaan media sosial oleh remaja, peserta didik dapat menentukan sikap muslim yang paling tepat.",
        "tingkat": "Rendah", "kunci": "B",
        "soal": (
            "Media sosial menjadi bagian tidak terpisahkan dari kehidupan remaja zaman sekarang. Sebagai seorang muslim, sikap yang paling tepat ketika menggunakan media sosial adalah ...\n"
            "A. Menggunakan sebebas-bebasnya tanpa batasan\n"
            "B. Menggunakan dengan bijak sesuai adab Islam dan menjadikannya sarana kebaikan\n"
            "C. Meninggalkan media sosial sepenuhnya karena diharamkan\n"
            "D. Hanya digunakan untuk mencari hiburan tanpa manfaat\n"
            "E. Digunakan untuk menjatuhkan orang lain"
        ),
    },
    {
        "no": 7, "bab": "Adab Menggunakan Media Sosial",
        "elemen": "Akhlak – Adab Menggunakan Media Sosial",
        "cp": "Peserta didik memahami konsep tabayyun berdasarkan QS Al-Hujurat ayat 6.",
        "tp": "Peserta didik dapat menjelaskan pengertian tabayyun.",
        "indikator_kompetensi": "Memahami konsep tabayyun.",
        "indikator_soal": "Disajikan ayat QS Al-Hujurat ayat 6, peserta didik dapat menentukan istilah yang sesuai dengan sikap memeriksa kebenaran berita.",
        "tingkat": "Sedang", "kunci": "C",
        "soal": (
            "Allah Swt. berfirman dalam QS Al-Hujurat ayat 6 yang memerintahkan orang beriman untuk memeriksa kebenaran sebuah berita sebelum menyebarkannya. Sikap memeriksa kebenaran informasi ini disebut ...\n"
            "A. Ghibah\n"
            "B. Namimah\n"
            "C. Tabayyun\n"
            "D. Su'udzon\n"
            "E. Fitnah"
        ),
    },
    {
        "no": 8, "bab": "Adab Menggunakan Media Sosial",
        "elemen": "Akhlak – Adab Menggunakan Media Sosial",
        "cp": "Peserta didik memahami bahaya hoax di media sosial.",
        "tp": "Peserta didik dapat mengategorikan hoax dalam pandangan Islam.",
        "indikator_kompetensi": "Mengategorikan perbuatan menyebar hoax.",
        "indikator_soal": "Disajikan deskripsi penyebaran berita bohong, peserta didik dapat menentukan kategori perbuatannya dalam Islam.",
        "tingkat": "Sedang", "kunci": "C",
        "soal": (
            "Salah satu dampak negatif media sosial yang banyak menjangkiti penggunanya adalah menyebarkan berita bohong atau berita palsu yang belum tentu kebenarannya. Berita bohong ini dalam Islam dikategorikan sebagai ...\n"
            "A. Akhlak terpuji\n"
            "B. Berita yang dianjurkan disebarkan\n"
            "C. Hoax dan termasuk perbuatan dosa karena mengandung kebohongan\n"
            "D. Hiburan yang menyenangkan\n"
            "E. Sarana dakwah yang baik"
        ),
    },
    {
        "no": 9, "bab": "Adab Menggunakan Media Sosial",
        "elemen": "Akhlak – Adab Menggunakan Media Sosial",
        "cp": "Peserta didik memahami manfaat positif media sosial bagi seorang muslim.",
        "tp": "Peserta didik dapat menentukan manfaat positif media sosial bagi seorang muslim.",
        "indikator_kompetensi": "Memahami manfaat positif media sosial.",
        "indikator_soal": "Disajikan pernyataan tentang manfaat sosmed, peserta didik dapat menentukan manfaat positif sesuai ajaran Islam.",
        "tingkat": "Sedang", "kunci": "B",
        "soal": (
            "Selain memiliki dampak negatif, media sosial juga memiliki manfaat positif jika digunakan dengan benar. Salah satu manfaat positif media sosial bagi seorang muslim adalah ...\n"
            "A. Untuk menyebarkan aib teman\n"
            "B. Menjadi sarana dakwah, menyebarkan ilmu, dan mempererat silaturahmi\n"
            "C. Untuk pamer kekayaan dan kemewahan\n"
            "D. Untuk membicarakan kekurangan orang lain\n"
            "E. Untuk menyaingi prestasi orang lain"
        ),
    },
    {
        "no": 10, "bab": "Adab Menggunakan Media Sosial",
        "elemen": "Akhlak – Adab Menggunakan Media Sosial",
        "cp": "Peserta didik memahami perbuatan tercela di media sosial.",
        "tp": "Peserta didik dapat menentukan istilah perbuatan membicarakan keburukan orang lain.",
        "indikator_kompetensi": "Mengidentifikasi ghibah di media sosial.",
        "indikator_soal": "Disajikan deskripsi perbuatan membicarakan keburukan orang lain meskipun benar, peserta didik dapat menentukan istilahnya.",
        "tingkat": "Sedang", "kunci": "A",
        "soal": (
            "Membicarakan keburukan orang lain meskipun benar adanya termasuk perbuatan tercela yang sering terjadi di media sosial. Perbuatan tersebut dalam Islam dikenal dengan istilah ...\n"
            "A. Ghibah\n"
            "B. Tabayyun\n"
            "C. Tasamuh\n"
            "D. Ta'aruf\n"
            "E. Tawadhu"
        ),
    },
    # ===== BAB 3: TOLERANSI DAN MEMELIHARA KEHIDUPAN MANUSIA (5 soal) =====
    {
        "no": 11, "bab": "Toleransi dan Memelihara Kehidupan Manusia",
        "elemen": "Akhlak – Toleransi dan Memelihara Kehidupan Manusia",
        "cp": "Peserta didik memahami pengertian toleransi (tasamuh).",
        "tp": "Peserta didik dapat menjelaskan pengertian toleransi.",
        "indikator_kompetensi": "Memahami pengertian toleransi.",
        "indikator_soal": "Disajikan beberapa pengertian tasamuh, peserta didik dapat menentukan pengertian yang paling tepat.",
        "tingkat": "Rendah", "kunci": "B",
        "soal": (
            "Toleransi dalam istilah Islam dikenal dengan tasamuh. Pengertian toleransi (tasamuh) yang paling tepat adalah ...\n"
            "A. Mengikuti seluruh ajaran agama lain\n"
            "B. Sikap saling menghargai, menghormati, dan tidak memaksakan kehendak kepada orang lain yang berbeda keyakinan\n"
            "C. Mencampuradukkan ibadah seluruh agama\n"
            "D. Menganggap semua agama sama benarnya\n"
            "E. Menjauhi orang yang berbeda agama"
        ),
    },
    {
        "no": 12, "bab": "Toleransi dan Memelihara Kehidupan Manusia",
        "elemen": "Akhlak – Toleransi dan Memelihara Kehidupan Manusia",
        "cp": "Peserta didik memahami dasar toleransi dari QS Al-Kafirun.",
        "tp": "Peserta didik dapat menjelaskan kandungan QS Al-Kafirun ayat 6.",
        "indikator_kompetensi": "Memahami kandungan QS Al-Kafirun ayat 6.",
        "indikator_soal": "Disajikan kutipan QS Al-Kafirun ayat 6, peserta didik dapat menentukan prinsip toleransi yang terkandung.",
        "tingkat": "Sedang", "kunci": "C",
        "soal": (
            "Allah Swt. berfirman dalam QS Al-Kafirun ayat 6: \"Lakum dinukum wa liya din\" yang artinya \"untukmu agamamu dan untukku agamaku\". Ayat tersebut menjadi dasar prinsip toleransi yaitu ...\n"
            "A. Boleh saling mencampuri urusan ibadah agama lain\n"
            "B. Wajib mengikuti agama mayoritas\n"
            "C. Saling menghormati keyakinan masing-masing tanpa mencampuradukkan akidah\n"
            "D. Larangan bergaul dengan non-muslim sama sekali\n"
            "E. Perintah memaksa orang lain masuk Islam"
        ),
    },
    {
        "no": 13, "bab": "Toleransi dan Memelihara Kehidupan Manusia",
        "elemen": "Akhlak – Toleransi dan Memelihara Kehidupan Manusia",
        "cp": "Peserta didik memahami batasan toleransi dalam Islam.",
        "tp": "Peserta didik dapat membedakan toleransi yang dibolehkan dan tidak dibolehkan.",
        "indikator_kompetensi": "Memahami batasan toleransi.",
        "indikator_soal": "Disajikan beberapa bentuk perilaku, peserta didik dapat menentukan bentuk toleransi yang TIDAK diperbolehkan.",
        "tingkat": "Sedang", "kunci": "D",
        "soal": (
            "Toleransi dalam Islam memiliki batasan yang jelas. Berikut ini bentuk toleransi yang TIDAK diperbolehkan dalam Islam adalah ...\n"
            "A. Menghormati tetangga yang berbeda agama\n"
            "B. Membantu korban bencana tanpa membedakan agama\n"
            "C. Hidup berdampingan secara damai dengan pemeluk agama lain\n"
            "D. Mengikuti tata cara ibadah agama lain atau mencampuradukkan akidah\n"
            "E. Menjaga kebersihan lingkungan bersama warga yang berbeda agama"
        ),
    },
    {
        "no": 14, "bab": "Toleransi dan Memelihara Kehidupan Manusia",
        "elemen": "Akhlak – Toleransi dan Memelihara Kehidupan Manusia",
        "cp": "Peserta didik memahami QS Al-Maidah ayat 32 tentang memelihara kehidupan.",
        "tp": "Peserta didik dapat menjelaskan kandungan QS Al-Maidah ayat 32.",
        "indikator_kompetensi": "Memahami kandungan QS Al-Maidah ayat 32.",
        "indikator_soal": "Disajikan kandungan QS Al-Maidah ayat 32, peserta didik dapat menentukan kesimpulan ajarannya.",
        "tingkat": "Sedang", "kunci": "B",
        "soal": (
            "Dalam QS Al-Maidah ayat 32 ditegaskan bahwa membunuh seorang manusia tanpa alasan yang dibenarkan agama sama dengan membunuh seluruh manusia. Ayat ini menunjukkan bahwa Islam sangat ...\n"
            "A. Membolehkan kekerasan\n"
            "B. Menjunjung tinggi nilai kehidupan dan memelihara jiwa manusia\n"
            "C. Mengabaikan nasib manusia\n"
            "D. Membiarkan permusuhan antar manusia\n"
            "E. Membatasi gerak manusia"
        ),
    },
    {
        "no": 15, "bab": "Toleransi dan Memelihara Kehidupan Manusia",
        "elemen": "Akhlak – Toleransi dan Memelihara Kehidupan Manusia",
        "cp": "Peserta didik dapat menerapkan toleransi dalam kehidupan sehari-hari.",
        "tp": "Peserta didik dapat menentukan contoh sikap toleransi yang tepat.",
        "indikator_kompetensi": "Menerapkan sikap toleransi.",
        "indikator_soal": "Disajikan ilustrasi sikap seorang siswa, peserta didik dapat menentukan kesesuaian dengan toleransi Islam.",
        "tingkat": "Sulit", "kunci": "B",
        "soal": (
            "Perhatikan pernyataan berikut: \"Andi tetap berteman baik dan saling membantu dengan Budi yang beragama Kristen, namun Andi tidak ikut menjalankan ibadah agama Budi.\" Sikap Andi mencerminkan ...\n"
            "A. Sinkretisme agama\n"
            "B. Sikap toleransi yang benar sesuai ajaran Islam\n"
            "C. Sikap fanatik buta\n"
            "D. Sikap acuh tak acuh\n"
            "E. Pelanggaran terhadap akidah Islam"
        ),
    },
    # ===== BAB 4: CABANG IMAN (5 soal) =====
    {
        "no": 16, "bab": "Cabang Iman: Kehormatan, Ikhlas, Malu, dan Zuhud",
        "elemen": "Akidah – Syu'abul Iman (Cabang Iman)",
        "cp": "Peserta didik memahami cabang-cabang iman (syu'abul iman).",
        "tp": "Peserta didik dapat menyebutkan jumlah cabang iman berdasarkan hadis.",
        "indikator_kompetensi": "Memahami jumlah cabang iman.",
        "indikator_soal": "Disajikan pernyataan tentang syu'abul iman, peserta didik dapat menentukan jumlah cabang iman berdasarkan hadis Nabi saw.",
        "tingkat": "Sedang", "kunci": "D",
        "soal": (
            "Dalam ajaran Islam, iman memiliki banyak cabang yang dikenal dengan istilah syu'abul iman. Berdasarkan hadis Rasulullah saw., cabang iman berjumlah ...\n"
            "A. 7 cabang\n"
            "B. 17 cabang\n"
            "C. 27 cabang\n"
            "D. lebih dari 70 cabang\n"
            "E. tidak terbatas"
        ),
    },
    {
        "no": 17, "bab": "Cabang Iman: Kehormatan, Ikhlas, Malu, dan Zuhud",
        "elemen": "Akhlak – Menjaga Kehormatan (Hifzul 'Irdh)",
        "cp": "Peserta didik memahami perilaku menjaga kehormatan diri.",
        "tp": "Peserta didik dapat memberikan contoh perilaku menjaga kehormatan.",
        "indikator_kompetensi": "Menerapkan perilaku menjaga kehormatan.",
        "indikator_soal": "Disajikan beberapa pilihan perilaku, peserta didik dapat menentukan contoh perilaku menjaga kehormatan diri.",
        "tingkat": "Sedang", "kunci": "A",
        "soal": (
            "Menjaga kehormatan diri (hifzul 'irdh) merupakan salah satu cabang iman yang sangat penting. Berikut ini contoh perilaku menjaga kehormatan diri adalah ...\n"
            "A. Berpakaian sopan menutup aurat, menjaga lisan, dan menjauhi pergaulan bebas\n"
            "B. Berpakaian seksi untuk menarik perhatian\n"
            "C. Sering membicarakan keburukan orang lain\n"
            "D. Mengikuti tren tanpa memperhatikan syariat\n"
            "E. Mencari sensasi di media sosial"
        ),
    },
    {
        "no": 18, "bab": "Cabang Iman: Kehormatan, Ikhlas, Malu, dan Zuhud",
        "elemen": "Akhlak – Ikhlas",
        "cp": "Peserta didik memahami pengertian ikhlas.",
        "tp": "Peserta didik dapat menjelaskan pengertian ikhlas.",
        "indikator_kompetensi": "Memahami pengertian ikhlas.",
        "indikator_soal": "Disajikan deskripsi tentang ibadah karena Allah, peserta didik dapat menentukan istilah yang sesuai.",
        "tingkat": "Rendah", "kunci": "C",
        "soal": (
            "Sikap melakukan ibadah maupun perbuatan baik semata-mata karena Allah Swt. tanpa mengharap pujian dari manusia disebut ...\n"
            "A. Riya\n"
            "B. Sum'ah\n"
            "C. Ikhlas\n"
            "D. Ujub\n"
            "E. Takabur"
        ),
    },
    {
        "no": 19, "bab": "Cabang Iman: Kehormatan, Ikhlas, Malu, dan Zuhud",
        "elemen": "Akhlak – Malu (Al-Haya')",
        "cp": "Peserta didik memahami pengertian malu (al-haya').",
        "tp": "Peserta didik dapat menjelaskan pengertian malu sebagai cabang iman.",
        "indikator_kompetensi": "Memahami pengertian malu.",
        "indikator_soal": "Disajikan hadis tentang malu, peserta didik dapat menentukan pengertian malu yang tepat dalam Islam.",
        "tingkat": "Sedang", "kunci": "C",
        "soal": (
            "Sifat malu merupakan salah satu cabang iman. Rasulullah saw. bersabda: \"Al-haya'u min al-iman\" (malu sebagian dari iman). Yang dimaksud malu (al-haya') dalam Islam adalah ...\n"
            "A. Malu mengerjakan kebaikan\n"
            "B. Malu untuk berkata jujur\n"
            "C. Sikap yang mendorong seseorang meninggalkan perbuatan buruk dan mengerjakan kewajiban\n"
            "D. Malu untuk beribadah di tempat umum\n"
            "E. Malu untuk menuntut ilmu"
        ),
    },
    {
        "no": 20, "bab": "Cabang Iman: Kehormatan, Ikhlas, Malu, dan Zuhud",
        "elemen": "Akhlak – Zuhud",
        "cp": "Peserta didik memahami pengertian zuhud.",
        "tp": "Peserta didik dapat menjelaskan pengertian zuhud yang benar.",
        "indikator_kompetensi": "Memahami pengertian zuhud.",
        "indikator_soal": "Disajikan beberapa pengertian zuhud, peserta didik dapat menentukan pengertian zuhud yang paling tepat.",
        "tingkat": "Sulit", "kunci": "C",
        "soal": (
            "Zuhud merupakan akhlak terpuji yang sangat dianjurkan dalam Islam. Pengertian zuhud yang paling tepat adalah ...\n"
            "A. Meninggalkan dunia sepenuhnya dan hidup sebagai pertapa\n"
            "B. Menolak segala bentuk harta benda\n"
            "C. Sikap tidak terikat pada kemewahan dunia, menggunakan dunia secukupnya, dan menjadikan akhirat sebagai tujuan utama\n"
            "D. Hidup dalam kemiskinan sengaja\n"
            "E. Tidak mau bekerja mencari nafkah"
        ),
    },
]

# Essay
ESSAY = [
    # ===== BAB 1: PERNIKAHAN (3 soal) =====
    {
        "no": 1, "bab": "Pernikahan dalam Islam",
        "elemen": "Fikih – Pernikahan dalam Islam",
        "cp": "Peserta didik memahami pengertian, hukum, dan hikmah pernikahan.",
        "tp": "Peserta didik dapat menjelaskan pengertian, hukum, dan hikmah pernikahan.",
        "indikator_kompetensi": "Menjelaskan pernikahan secara komprehensif.",
        "indikator_soal": "Disajikan narasi tentang pernikahan, peserta didik dapat menjelaskan pengertian, 5 hukum, dan 3 hikmah pernikahan.",
        "tingkat": "Sedang",
        "soal": (
            "Pernikahan adalah ibadah yang sangat dianjurkan dalam Islam. Dalam QS An-Nur ayat 32 Allah Swt. memerintahkan untuk menikahkan orang-orang yang sendirian (lajang) dari kalangan kalian. Pernikahan bukan sekadar mempertemukan dua insan, tetapi membangun keluarga yang sakinah, mawaddah, dan rahmah sesuai tuntunan syariat.\n"
            "Jelaskan pengertian nikah dalam Islam, sebutkan hukum nikah berdasarkan kondisi seseorang (5 hukum), dan tuliskan 3 hikmah dari pernikahan!"
        ),
        "jawaban": (
            "Pengertian nikah: ikatan lahir batin antara seorang laki-laki dan perempuan sebagai suami istri yang sah berdasarkan syariat Islam, untuk membentuk keluarga sakinah, mawaddah, wa rahmah.\n"
            "\n"
            "Hukum nikah (5):\n"
            "1) Wajib – bagi yang sudah mampu lahir batin dan dikhawatirkan terjerumus zina jika tidak menikah.\n"
            "2) Sunah – bagi yang sudah mampu dan tetap mampu menahan diri dari zina.\n"
            "3) Mubah – bagi yang mampu tetapi belum ada dorongan kuat menikah.\n"
            "4) Makruh – bagi yang belum mampu memberi nafkah lahir/batin.\n"
            "5) Haram – bagi yang menikahi dengan niat menyakiti, menipu, atau menyengsarakan pasangan.\n"
            "\n"
            "Hikmah pernikahan (3):\n"
            "1) Menjaga kehormatan diri dan menghindarkan dari perbuatan zina.\n"
            "2) Melahirkan keturunan yang sah dan saleh-salehah.\n"
            "3) Membentuk keluarga sakinah, mawaddah, wa rahmah sebagai sarana ketenangan hati dan ibadah kepada Allah Swt."
        ),
    },
    {
        "no": 2, "bab": "Pernikahan dalam Islam",
        "elemen": "Fikih – Pernikahan dalam Islam",
        "cp": "Peserta didik memahami rukun nikah.",
        "tp": "Peserta didik dapat menyebutkan dan menjelaskan rukun nikah.",
        "indikator_kompetensi": "Menjelaskan rukun nikah.",
        "indikator_soal": "Disajikan narasi tentang sahnya pernikahan, peserta didik dapat menyebutkan dan menjelaskan 5 rukun nikah.",
        "tingkat": "Sedang",
        "soal": (
            "Pernikahan dianggap sah apabila telah memenuhi seluruh rukun dan syaratnya. Rukun nikah adalah sesuatu yang harus ada saat akad nikah berlangsung; jika ada yang tidak terpenuhi, pernikahan menjadi tidak sah.\n"
            "Sebutkan dan jelaskan rukun-rukun nikah dalam Islam!"
        ),
        "jawaban": (
            "Rukun nikah ada 5:\n"
            "1) Calon suami – laki-laki muslim, baligh, berakal, bukan mahram dari calon istri, dan ridha menikah.\n"
            "2) Calon istri – perempuan muslimah, bukan mahram, tidak sedang dalam masa iddah, dan ridha dinikahi.\n"
            "3) Wali nikah – wali dari pihak perempuan (urutan: ayah, kakek, saudara laki-laki, dst.). Tanpa wali, nikah tidak sah.\n"
            "4) Dua orang saksi – laki-laki muslim, baligh, berakal, dan adil.\n"
            "5) Ijab dan kabul (sigat) – ucapan menyerahkan dari wali dan ucapan menerima dari mempelai laki-laki yang dilakukan dengan jelas, langsung, dan saling berkaitan."
        ),
    },
    {
        "no": 3, "bab": "Pernikahan dalam Islam",
        "elemen": "Fikih – Pernikahan dalam Islam",
        "cp": "Peserta didik memahami wanita yang haram dinikahi (mahram).",
        "tp": "Peserta didik dapat menyebutkan wanita yang haram dinikahi.",
        "indikator_kompetensi": "Mengidentifikasi mahram.",
        "indikator_soal": "Disajikan narasi tentang aturan mahram, peserta didik dapat menyebutkan minimal 4 mahram nasab dan 2 mahram musaharah.",
        "tingkat": "Sulit",
        "soal": (
            "Dalam Islam terdapat aturan tentang siapa saja yang haram dinikahi (mahram), baik karena hubungan nasab, persusuan, maupun pernikahan. Hal ini bertujuan menjaga kesucian hubungan keluarga dan kemurnian keturunan.\n"
            "Sebutkan minimal 4 wanita yang haram dinikahi karena hubungan nasab dan minimal 2 wanita yang haram dinikahi karena hubungan pernikahan (musaharah)!"
        ),
        "jawaban": (
            "Haram dinikahi karena nasab (minimal 4):\n"
            "1) Ibu dan ibu ke atas (nenek, dst.).\n"
            "2) Anak perempuan dan ke bawah (cucu perempuan).\n"
            "3) Saudara perempuan (kandung/seayah/seibu).\n"
            "4) Bibi dari pihak ayah ('ammah).\n"
            "5) Bibi dari pihak ibu (khalah).\n"
            "6) Anak perempuan dari saudara laki-laki (keponakan).\n"
            "7) Anak perempuan dari saudara perempuan (keponakan).\n"
            "\n"
            "Haram dinikahi karena hubungan pernikahan/musaharah (minimal 2):\n"
            "1) Mertua perempuan (ibu dari istri).\n"
            "2) Anak tiri perempuan (anak istri dari suami sebelumnya) yang ibunya telah dicampuri.\n"
            "3) Menantu perempuan (istri dari anak laki-laki kandung).\n"
            "4) Ibu tiri (istri dari ayah kandung)."
        ),
    },
    # ===== BAB 2: ADAB MENGGUNAKAN MEDIA SOSIAL (2 soal) =====
    {
        "no": 4, "bab": "Adab Menggunakan Media Sosial",
        "elemen": "Akhlak – Adab Menggunakan Media Sosial",
        "cp": "Peserta didik memahami adab muslim dalam menggunakan media sosial.",
        "tp": "Peserta didik dapat menyebutkan adab bermedia sosial.",
        "indikator_kompetensi": "Menyebutkan adab bermedia sosial.",
        "indikator_soal": "Disajikan narasi tentang pelanggaran adab di sosmed, peserta didik dapat menyebutkan minimal 5 adab muslim dalam bermedia sosial.",
        "tingkat": "Sedang",
        "soal": (
            "Media sosial telah menjadi bagian penting dari kehidupan remaja modern. Akan tetapi, banyak penggunaan media sosial yang justru melanggar adab dan akhlak Islam, seperti mengumbar aib, menyebarkan hoax, atau melakukan perundungan (cyber bullying).\n"
            "Sebutkan minimal 5 adab seorang muslim dalam menggunakan media sosial!"
        ),
        "jawaban": (
            "Adab muslim dalam bermedia sosial (minimal 5):\n"
            "1) Tabayyun – memeriksa kebenaran informasi sebelum mempercayai/menyebarkannya.\n"
            "2) Menjaga lisan dan jari – tidak menulis/menyebarkan ujaran kebencian, ghibah, fitnah, atau cacian.\n"
            "3) Menjaga privasi diri dan orang lain – tidak mengumbar aib pribadi maupun orang lain.\n"
            "4) Menggunakan untuk hal positif – belajar, dakwah, silaturahmi, dan berbagi kebaikan.\n"
            "5) Tidak berlebihan – menjaga keseimbangan dengan ibadah, belajar, dan keluarga.\n"
            "6) Tidak menampilkan konten yang melanggar syariat (membuka aurat, vulgar, pamer kemewahan, dll.).\n"
            "7) Menebar manfaat dengan konten inspiratif, dakwah, dan mendidik."
        ),
    },
    {
        "no": 5, "bab": "Adab Menggunakan Media Sosial",
        "elemen": "Akhlak – Adab Menggunakan Media Sosial",
        "cp": "Peserta didik memahami konsep tabayyun dalam menghadapi hoax.",
        "tp": "Peserta didik dapat menjelaskan tabayyun dan langkah penerapannya.",
        "indikator_kompetensi": "Menjelaskan tabayyun dan penerapannya.",
        "indikator_soal": "Disajikan narasi tentang fenomena hoax dan QS Al-Hujurat ayat 6, peserta didik dapat menjelaskan pengertian, pentingnya, dan langkah tabayyun.",
        "tingkat": "Sedang",
        "soal": (
            "Salah satu fenomena yang sering terjadi di media sosial adalah penyebaran berita bohong (hoax). Allah Swt. berfirman dalam QS Al-Hujurat ayat 6 yang memerintahkan orang beriman untuk melakukan tabayyun terlebih dahulu sebelum menerima sebuah berita.\n"
            "Jelaskan pengertian tabayyun, mengapa tabayyun penting dalam bermedia sosial, dan sebutkan langkah-langkah tabayyun yang bisa dilakukan!"
        ),
        "jawaban": (
            "Pengertian tabayyun: meneliti, memeriksa, dan mengklarifikasi kebenaran sebuah berita/informasi sebelum dipercaya atau disebarkan.\n"
            "\n"
            "Pentingnya tabayyun di media sosial:\n"
            "• Mencegah penyebaran hoax/fitnah yang merugikan orang lain.\n"
            "• Menghindari dosa karena menyebar kebohongan.\n"
            "• Menjaga keharmonisan masyarakat dan menghindari konflik.\n"
            "• Menjaga akidah agar tidak mudah diadu domba.\n"
            "\n"
            "Langkah-langkah tabayyun:\n"
            "1) Memeriksa sumber berita – apakah berasal dari media/akun yang terpercaya.\n"
            "2) Cross-check ke media kredibel lain yang relevan.\n"
            "3) Memastikan konteks, waktu, dan tempat berita.\n"
            "4) Tidak terburu-buru menyebarkan sebelum yakin akan kebenarannya.\n"
            "5) Bila masih ragu, sebaiknya tidak disebarkan."
        ),
    },
    # ===== BAB 3: TOLERANSI DAN MEMELIHARA KEHIDUPAN MANUSIA (2 soal) =====
    {
        "no": 6, "bab": "Toleransi dan Memelihara Kehidupan Manusia",
        "elemen": "Akhlak – Toleransi (Tasamuh)",
        "cp": "Peserta didik memahami konsep dan batasan toleransi.",
        "tp": "Peserta didik dapat menjelaskan pengertian, batasan, dan penerapan toleransi.",
        "indikator_kompetensi": "Menjelaskan pengertian, batasan, dan contoh toleransi.",
        "indikator_soal": "Disajikan narasi tentang toleransi di Indonesia, peserta didik dapat menjelaskan pengertian, batasan, dan minimal 3 contoh yang dibolehkan & 2 yang tidak dibolehkan.",
        "tingkat": "Sulit",
        "soal": (
            "Toleransi (tasamuh) adalah salah satu nilai mulia dalam ajaran Islam. Indonesia sebagai negara dengan beragam suku, agama, dan budaya sangat membutuhkan sikap toleransi agar terwujud kerukunan. Namun, toleransi dalam Islam memiliki batasan yang jelas, yaitu tidak boleh sampai mencampuradukkan akidah.\n"
            "Jelaskan pengertian toleransi (tasamuh) dalam Islam, batasannya, dan berikan minimal 3 contoh penerapan toleransi yang DIBOLEHKAN serta 2 contoh yang TIDAK DIBOLEHKAN!"
        ),
        "jawaban": (
            "Pengertian toleransi (tasamuh): sikap saling menghargai dan menghormati keyakinan, pendapat, dan tata cara hidup orang lain tanpa memaksakan kehendak, selama tidak melanggar prinsip akidah dan syariat.\n"
            "\n"
            "Batasan: toleransi dibolehkan dalam urusan muamalah/sosial, tidak dibolehkan dalam urusan akidah dan ibadah (\"Lakum dinukum waliyadin\" – QS Al-Kafirun: 6).\n"
            "\n"
            "Contoh toleransi yang DIBOLEHKAN (minimal 3):\n"
            "1) Menghormati tetangga yang berbeda agama.\n"
            "2) Menjenguk yang sakit dan membantu korban bencana tanpa membedakan agama.\n"
            "3) Bekerja sama dalam hal kebaikan (gotong royong, kerja bakti, jaga kebersihan).\n"
            "4) Tidak mengganggu pemeluk agama lain saat menjalankan ibadahnya.\n"
            "\n"
            "Contoh yang TIDAK DIBOLEHKAN (minimal 2):\n"
            "1) Mengikuti ritual ibadah agama lain.\n"
            "2) Meyakini bahwa semua agama itu sama benarnya (pluralisme akidah).\n"
            "3) Mencampuradukkan ajaran agama (sinkretisme)."
        ),
    },
    {
        "no": 7, "bab": "Toleransi dan Memelihara Kehidupan Manusia",
        "elemen": "Akhlak – Memelihara Kehidupan Manusia",
        "cp": "Peserta didik memahami kandungan QS Al-Maidah ayat 32.",
        "tp": "Peserta didik dapat menjelaskan kandungan QS Al-Maidah ayat 32 dan penerapannya.",
        "indikator_kompetensi": "Menjelaskan kandungan QS Al-Maidah ayat 32.",
        "indikator_soal": "Disajikan narasi tentang QS Al-Maidah ayat 32, peserta didik dapat menjelaskan kandungannya dan menyebutkan minimal 3 perilaku memelihara kehidupan.",
        "tingkat": "Sedang",
        "soal": (
            "QS Al-Maidah ayat 32 menjelaskan tentang betapa agungnya nilai sebuah jiwa manusia. Membunuh satu jiwa tanpa alasan yang dibenarkan, seakan-akan ia telah membunuh seluruh manusia, dan barangsiapa memelihara kehidupan seorang manusia, seakan-akan ia telah memelihara kehidupan seluruh manusia.\n"
            "Jelaskan isi kandungan QS Al-Maidah ayat 32 dan sebutkan minimal 3 bentuk perilaku memelihara kehidupan manusia dalam kehidupan sehari-hari!"
        ),
        "jawaban": (
            "Kandungan QS Al-Maidah ayat 32:\n"
            "• Islam sangat menjunjung tinggi dan memuliakan nilai jiwa/kehidupan manusia.\n"
            "• Membunuh satu jiwa tanpa hak (qishash atau alasan syar'i) merupakan dosa besar dan diibaratkan membunuh seluruh manusia.\n"
            "• Sebaliknya, memelihara/menyelamatkan satu jiwa nilainya seperti menyelamatkan seluruh manusia.\n"
            "• Ayat ini menjadi dasar prinsip hifzh an-nafs (menjaga jiwa) dalam maqashid syariah.\n"
            "\n"
            "Bentuk perilaku memelihara kehidupan (minimal 3):\n"
            "1) Tidak menyakiti, melukai, apalagi membunuh orang lain.\n"
            "2) Menolong korban kecelakaan, mendonorkan darah, membantu kaum fakir/miskin.\n"
            "3) Menjaga lingkungan, menjauhi tawuran/perundungan/kekerasan.\n"
            "4) Mengonsumsi makanan halal dan sehat agar tubuh terpelihara.\n"
            "5) Menjauhi narkoba dan minuman keras yang merusak diri.\n"
            "6) Tidak menyebar hoax/provokasi yang dapat menimbulkan korban jiwa."
        ),
    },
    # ===== BAB 4: CABANG IMAN (3 soal) =====
    {
        "no": 8, "bab": "Cabang Iman: Kehormatan, Ikhlas, Malu, dan Zuhud",
        "elemen": "Akidah – Syu'abul Iman",
        "cp": "Peserta didik memahami konsep dan cabang-cabang iman.",
        "tp": "Peserta didik dapat menjelaskan syu'abul iman dan menyebutkan cabangnya.",
        "indikator_kompetensi": "Menjelaskan syu'abul iman.",
        "indikator_soal": "Disajikan narasi tentang iman dan amal, peserta didik dapat menjelaskan pengertian, jumlah, dan minimal 4 cabang iman.",
        "tingkat": "Sedang",
        "soal": (
            "Iman bukan hanya pengakuan dengan lisan dan keyakinan dalam hati, tetapi juga harus dibuktikan dengan amal perbuatan. Rasulullah saw. bersabda bahwa iman memiliki banyak cabang. Cabang-cabang iman inilah yang dikenal dengan istilah syu'abul iman.\n"
            "Jelaskan pengertian syu'abul iman, jumlah cabang iman berdasarkan hadis Nabi saw., dan sebutkan minimal 4 cabang iman beserta penjelasannya!"
        ),
        "jawaban": (
            "Pengertian syu'abul iman: cabang-cabang iman, yaitu berbagai bentuk perilaku, ucapan, dan keyakinan yang merupakan manifestasi dari keimanan seorang muslim.\n"
            "\n"
            "Jumlah cabang iman: Berdasarkan hadis Nabi saw. (HR Muslim), \"Iman itu memiliki lebih dari 70 cabang. Yang paling utama adalah ucapan 'La ilaha illallah' dan yang paling rendah adalah menyingkirkan gangguan/duri dari jalan, dan malu adalah sebagian dari iman.\" Dalam riwayat lain disebutkan 77 cabang.\n"
            "\n"
            "Contoh cabang iman (minimal 4):\n"
            "1) Mengucapkan dan meyakini kalimat tauhid (La ilaha illallah) – cabang iman tertinggi.\n"
            "2) Menyingkirkan duri/gangguan dari jalan – cabang iman terendah.\n"
            "3) Malu (al-haya') – sebagian dari iman yang mendorong meninggalkan dosa.\n"
            "4) Menjaga kehormatan (hifzul 'irdh) – menutup aurat, menjaga lisan, menjauhi zina.\n"
            "5) Ikhlas dalam beramal.\n"
            "6) Mencintai sesama muslim seperti mencintai diri sendiri.\n"
            "7) Sabar saat menghadapi cobaan dan syukur saat mendapat nikmat."
        ),
    },
    {
        "no": 9, "bab": "Cabang Iman: Kehormatan, Ikhlas, Malu, dan Zuhud",
        "elemen": "Akhlak – Ikhlas, Malu, dan Zuhud",
        "cp": "Peserta didik memahami ikhlas, malu, dan zuhud serta penerapannya.",
        "tp": "Peserta didik dapat menjelaskan ikhlas, malu, zuhud dan memberi contohnya.",
        "indikator_kompetensi": "Menjelaskan ikhlas, malu, dan zuhud.",
        "indikator_soal": "Disajikan narasi tentang akhlak terpuji, peserta didik dapat menjelaskan ikhlas, malu, dan zuhud beserta contoh penerapannya.",
        "tingkat": "Sedang",
        "soal": (
            "Akhlak terpuji yang sangat dianjurkan dalam Islam antara lain ikhlas, malu, dan zuhud. Ketiga sifat ini saling melengkapi dalam membentuk pribadi muslim yang mulia.\n"
            "Jelaskan pengertian ikhlas, malu (al-haya'), dan zuhud, serta berikan masing-masing satu contoh penerapannya dalam kehidupan sehari-hari!"
        ),
        "jawaban": (
            "1) IKHLAS\n"
            "   Pengertian: melakukan ibadah/amal kebaikan semata-mata karena Allah Swt. tanpa mengharap pujian, balasan, atau penghargaan dari manusia.\n"
            "   Contoh: bersedekah secara diam-diam tanpa diceritakan ke orang lain agar tidak riya.\n"
            "\n"
            "2) MALU (AL-HAYA')\n"
            "   Pengertian: sikap yang mendorong seseorang untuk meninggalkan perbuatan tercela dan mendorongnya untuk melaksanakan kewajiban serta amal kebaikan.\n"
            "   Contoh: malu memakai pakaian yang membuka aurat, malu berbohong kepada orang tua atau guru.\n"
            "\n"
            "3) ZUHUD\n"
            "   Pengertian: sikap tidak terikat hatinya kepada gemerlap dunia, menggunakan dunia secukupnya, dan menjadikan akhirat sebagai tujuan utama – meskipun memiliki harta banyak.\n"
            "   Contoh: hidup sederhana meski memiliki rezeki cukup, lebih banyak menabung untuk sedekah daripada bermewah-mewah."
        ),
    },
    {
        "no": 10, "bab": "Cabang Iman: Kehormatan, Ikhlas, Malu, dan Zuhud",
        "elemen": "Akhlak – Penerapan Akhlak Terpuji",
        "cp": "Peserta didik mampu menerapkan akhlak terpuji dalam era media sosial.",
        "tp": "Peserta didik dapat memberikan contoh penerapan menjaga kehormatan, ikhlas, malu, dan zuhud.",
        "indikator_kompetensi": "Menerapkan akhlak terpuji di era media sosial.",
        "indikator_soal": "Disajikan narasi tentang tantangan akhlak di era sosmed, peserta didik dapat memberikan masing-masing 2 contoh penerapan menjaga kehormatan, ikhlas, malu, dan zuhud.",
        "tingkat": "Sulit",
        "soal": (
            "Di era globalisasi dan media sosial saat ini, akhlak terpuji seperti menjaga kehormatan, ikhlas, malu, dan zuhud sering terkikis oleh budaya pamer (flexing), riya, dan hedonisme. Sebagai pelajar muslim, kita ditantang untuk tetap berpegang teguh pada akhlak Islam.\n"
            "Tuliskan masing-masing 2 contoh penerapan menjaga kehormatan, ikhlas, malu, dan zuhud bagi pelajar muslim di era media sosial!"
        ),
        "jawaban": (
            "A. MENJAGA KEHORMATAN (HIFZUL 'IRDH):\n"
            "   1) Berpakaian sopan menutup aurat baik di dunia nyata maupun di foto/video media sosial.\n"
            "   2) Tidak menyebar foto/cerita pribadi yang dapat menjatuhkan kehormatan diri/keluarga.\n"
            "\n"
            "B. IKHLAS:\n"
            "   1) Belajar dan menuntut ilmu karena Allah Swt., bukan sekadar mengejar pujian/peringkat.\n"
            "   2) Berbagi konten dakwah/kebaikan di sosmed tanpa mengharap pujian (like/komentar).\n"
            "\n"
            "C. MALU (AL-HAYA'):\n"
            "   1) Malu mengunggah konten yang melanggar syariat (membuka aurat, kasar, vulgar).\n"
            "   2) Malu meninggalkan salat atau ibadah wajib lainnya meski sedang asyik bermain sosmed.\n"
            "\n"
            "D. ZUHUD:\n"
            "   1) Tidak ikut tren flexing/pamer barang mewah di media sosial.\n"
            "   2) Hidup sederhana dan mendahulukan kebutuhan daripada keinginan dunia."
        ),
    },
]

# Kisi-kisi (PG: 4 baris, sesuai 4 bab; Essay: 5 baris, kita gabung sesuai bab)
pg_kisi = [
    {"no": "1", "elemen": "Fikih – Pernikahan dalam Islam",
     "cp": "Memahami ketentuan dan hikmah pernikahan dalam Islam.",
     "tp": "Menjelaskan pengertian, hukum, rukun, wali, dan hikmah pernikahan.",
     "jml": "5", "no_soal": "1 - 5", "bentuk": "PG", "tingkat": "Rendah/Sedang"},
    {"no": "2", "elemen": "Akhlak – Adab Menggunakan Media Sosial",
     "cp": "Memahami adab seorang muslim dalam bermedia sosial.",
     "tp": "Menjelaskan sikap muslim, tabayyun, bahaya hoax/ghibah, dan manfaat positif media sosial.",
     "jml": "5", "no_soal": "6 - 10", "bentuk": "PG", "tingkat": "Rendah/Sedang"},
    {"no": "3", "elemen": "Akhlak – Toleransi dan Memelihara Kehidupan Manusia",
     "cp": "Memahami toleransi dan menghargai nilai kehidupan manusia.",
     "tp": "Menjelaskan tasamuh, QS Al-Kafirun ayat 6, batasan toleransi, QS Al-Maidah ayat 32, dan contoh penerapannya.",
     "jml": "5", "no_soal": "11 - 15", "bentuk": "PG", "tingkat": "Sedang/Sulit"},
    {"no": "4", "elemen": "Akidah & Akhlak – Cabang Iman",
     "cp": "Memahami cabang iman: menjaga kehormatan, ikhlas, malu, dan zuhud.",
     "tp": "Menjelaskan syu'abul iman, menjaga kehormatan, ikhlas, malu, dan zuhud.",
     "jml": "5", "no_soal": "16 - 20", "bentuk": "PG", "tingkat": "Sedang/Sulit"},
]

essay_kisi = [
    {"no": "1", "elemen": "Fikih – Pernikahan dalam Islam",
     "cp": "Memahami pengertian, hukum, hikmah, rukun, dan mahram pernikahan.",
     "tp": "Menjelaskan pernikahan secara komprehensif (definisi, hukum, hikmah, rukun, mahram).",
     "jml": "3", "no_soal": "1 - 3", "bentuk": "Esay", "tingkat": "Sedang/Sulit"},
    {"no": "2", "elemen": "Akhlak – Adab Menggunakan Media Sosial",
     "cp": "Memahami adab dan tabayyun dalam bermedia sosial.",
     "tp": "Menyebutkan adab bermedia sosial dan menjelaskan tabayyun beserta langkahnya.",
     "jml": "2", "no_soal": "4 - 5", "bentuk": "Esay", "tingkat": "Sedang"},
    {"no": "3", "elemen": "Akhlak – Toleransi (Tasamuh)",
     "cp": "Memahami konsep, batasan, dan penerapan toleransi.",
     "tp": "Menjelaskan tasamuh, batasannya, serta contoh yang dibolehkan dan tidak.",
     "jml": "1", "no_soal": "6", "bentuk": "Esay", "tingkat": "Sulit"},
    {"no": "4", "elemen": "Akhlak – Memelihara Kehidupan Manusia",
     "cp": "Memahami kandungan QS Al-Maidah ayat 32.",
     "tp": "Menjelaskan kandungan ayat dan menyebutkan perilaku memelihara kehidupan.",
     "jml": "1", "no_soal": "7", "bentuk": "Esay", "tingkat": "Sedang"},
    {"no": "5", "elemen": "Akidah & Akhlak – Cabang Iman & Penerapannya",
     "cp": "Memahami syu'abul iman dan menerapkan akhlak terpuji.",
     "tp": "Menjelaskan syu'abul iman, ikhlas, malu, zuhud, dan menerapkannya di era sosmed.",
     "jml": "3", "no_soal": "8 - 10", "bentuk": "Esay", "tingkat": "Sedang/Sulit"},
]


# =====================================================================
# HELPERS
# =====================================================================

def set_cell_text(cell, lines, bold=False, size=None):
    if isinstance(lines, str):
        lines = lines.split('\n')
    tc = cell._tc
    for p in tc.findall(qn('w:p')):
        tc.remove(p)
    for line in lines:
        p = cell.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        spacing = pPr.find(qn('w:spacing'))
        from docx.oxml import OxmlElement
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            pPr.append(spacing)
        spacing.set(qn('w:after'), '0')
        spacing.set(qn('w:line'), '240')
        spacing.set(qn('w:lineRule'), 'auto')
        run = p.add_run(line if line else '')
        run.font.name = 'Calibri'
        if size:
            run.font.size = docx.shared.Pt(size)
        if bold:
            run.bold = True


def set_question_cell(cell, soal_text):
    tc = cell._tc
    for p in tc.findall(qn('w:p')):
        tc.remove(p)
    from docx.oxml import OxmlElement
    for line in soal_text.split('\n'):
        p = cell.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        spacing = pPr.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            pPr.append(spacing)
        spacing.set(qn('w:after'), '0')
        spacing.set(qn('w:line'), '276')
        spacing.set(qn('w:lineRule'), 'auto')
        run = p.add_run(line if line else '')
        run.font.name = 'Calibri'
        run.font.size = docx.shared.Pt(11)


def fill_header_para(paragraph, prefix, value):
    p_el = paragraph._p
    for r in p_el.findall(qn('w:r')):
        p_el.remove(r)
    from docx.oxml import OxmlElement
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Calibri')
    rFonts.set(qn('w:hAnsi'), 'Calibri')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = f"{prefix} : {value}"
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    p_el.append(r)


# =====================================================================
# 1. Fill cover page paragraphs
# =====================================================================
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("Nama  Guru") or txt.startswith("Nama Guru"):
        fill_header_para(p, "Nama Guru", NAMA_GURU)
    elif txt.startswith("Mata Pelajaran"):
        fill_header_para(p, "Mata Pelajaran", f"{MAPEL} / Kelas {KELAS}")
    elif txt.startswith("Kelas"):
        fill_header_para(p, "Kelas", KELAS)


# =====================================================================
# 2. Fill Table 0 (intro)
# =====================================================================
t0 = doc.tables[0]
cell0 = t0.rows[0].cells[0]
intro_lines = [
    f"Nama Sekolah\t    : SMK 1 PANCASILA",
    f"Nama Penulis Soal\t: {NAMA_GURU}",
    f"Mata Pelajaran/Kelas\t: {MAPEL} / Kelas {KELAS}",
    f"Satuan Unit Kerja\t: SMK 1 PANCASILA AMBULU",
    f"Kurikulum\t\t: Kurikulum Merdeka",
    f"Jumlah Soal\t\t: 30 butir (20 PG + 10 Esay)",
]
set_cell_text(cell0, intro_lines, size=11)


# =====================================================================
# 3. Fill kisi-kisi table (Table 1)
# =====================================================================
t1 = doc.tables[1]

def fill_kisi_row(row, entry):
    cells = row.cells
    set_cell_text(cells[0], entry["no"], size=10)
    set_cell_text(cells[1], entry["elemen"], size=10)
    set_cell_text(cells[2], entry["cp"], size=10)
    set_cell_text(cells[3], entry["tp"], size=10)
    set_cell_text(cells[5], entry["jml"], size=10)
    set_cell_text(cells[6], entry["no_soal"], size=10)
    set_cell_text(cells[7], entry["bentuk"], size=10)
    set_cell_text(cells[8], entry["tingkat"], size=10)

for i, entry in enumerate(pg_kisi):
    fill_kisi_row(t1.rows[1 + i], entry)
for i, entry in enumerate(essay_kisi):
    fill_kisi_row(t1.rows[6 + i], entry)
set_cell_text(t1.rows[11].cells[5], "30", size=10, bold=True)
set_cell_text(t1.rows[11].cells[6], "PG: 1-20\nEsay: 1-10", size=10, bold=True)


# =====================================================================
# 4. Duplicate PG kartu soal (13 -> 20, need 7 more)
# =====================================================================
elements = list(body)
last_pg_tbl = elements[61]
last_pg_para = elements[62]
insert_after = last_pg_para
for _ in range(7):
    new_tbl = deepcopy(last_pg_tbl)
    new_para = deepcopy(last_pg_para)
    insert_after.addnext(new_tbl)
    new_tbl.addnext(new_para)
    insert_after = new_para


# =====================================================================
# 5. Fill kartu soal
# =====================================================================
def fill_kartu_soal(table, item, is_essay=False):
    rows = table.rows
    set_cell_text(rows[2].cells[0], item["elemen"], size=10)
    set_cell_text(rows[6].cells[0], item["cp"], size=10)
    set_cell_text(rows[10].cells[0], item["indikator_kompetensi"], size=10)
    set_cell_text(rows[12].cells[0], item["indikator_soal"], size=10)
    set_cell_text(rows[13].cells[0], [
        "Tingkat Kesukaran:",
        f"☑ {item['tingkat']}",
    ], size=10)
    set_cell_text(rows[5].cells[2], str(item["no"]), size=12, bold=True)
    if not is_essay:
        set_cell_text(rows[11].cells[2], [
            "KUNCI JAWABAN",
            "",
            item["kunci"],
        ], size=12, bold=True)
    else:
        set_cell_text(rows[11].cells[2], [
            "PEDOMAN JAWABAN",
            "(lihat kunci di bawah)",
        ], size=10, bold=True)

    soal_lines = item["soal"].split('\n')
    if is_essay:
        content = list(soal_lines)
        content.append("")
        content.append("Kunci Jawaban / Pedoman Penskoran:")
        content.extend(item["jawaban"].split('\n'))
    else:
        content = soal_lines
    set_question_cell(rows[3].cells[3], '\n'.join(content))

# Fill PG (tables 2..21)
for i, item in enumerate(PG):
    fill_kartu_soal(doc.tables[2 + i], item, is_essay=False)


# =====================================================================
# 6. Delete excess essay tables (21 -> 10, delete 11)
# =====================================================================
elements = list(body)
tbl_indexes = [i for i, e in enumerate(elements) if e.tag.split('}')[-1] == 'tbl']
to_delete_tbl_idx = tbl_indexes[32:]
to_remove = []
for ti in to_delete_tbl_idx:
    tbl = elements[ti]
    to_remove.append(tbl)
    if ti + 1 < len(elements):
        next_el = elements[ti + 1]
        if next_el.tag.split('}')[-1] == 'p':
            to_remove.append(next_el)
final_keep = []
for el in to_remove:
    has_sect = el.find('.//' + qn('w:sectPr')) is not None
    if has_sect:
        continue
    final_keep.append(el)
for el in final_keep:
    el.getparent().remove(el)


# =====================================================================
# 7. Fill essay (tables 22..31)
# =====================================================================
for i, item in enumerate(ESSAY):
    fill_kartu_soal(doc.tables[22 + i], item, is_essay=True)


# =====================================================================
# 8. Save
# =====================================================================
import os
os.makedirs('/home/user/dianindri/output', exist_ok=True)
output_path = '/home/user/dianindri/output/KARTU_SOAL_PAI_Kelas_XI_Genap_2025_2026.docx'
doc.save(output_path)
print(f"Saved to: {output_path}")
print(f"Total tables: {len(doc.tables)}")
